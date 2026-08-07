from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from conftest import make_invoice, text
from openpyxl import load_workbook

from invoice_extract.config import DEFAULT_COLUMNS, KNOWN_COLUMNS, Config, ConfigError
from invoice_extract.documents import UnsupportedDocument, discover
from invoice_extract.validate import build
from invoice_extract.workbook import field_value, write

CONFIG = Config()


# --------------------------------------------------------------------------
# config
# --------------------------------------------------------------------------


def test_missing_config_file_is_reported(tmp_path):
    with pytest.raises(ConfigError, match="not found"):
        Config.load(tmp_path / "nope.yaml")


def test_unknown_column_field_is_rejected(tmp_path):
    path = tmp_path / "c.yaml"
    path.write_text("columns:\n  vendor_naem: Vendor\n")
    with pytest.raises(ConfigError, match="unknown column"):
        Config.load(path)


def test_column_order_and_names_are_honoured(tmp_path):
    path = tmp_path / "c.yaml"
    path.write_text(
        "columns:\n"
        "  total_amount: Amt\n"
        "  vendor_name: Supplier\n"
        "  gl_account: GL\n"
    )
    config = Config.load(path)
    assert list(config.columns) == ["total_amount", "vendor_name", "gl_account"]
    assert config.columns["vendor_name"] == "Supplier"


def test_every_configurable_column_can_be_exported():
    """A column the config accepts but the exporter cannot produce would fail
    mid-run, after the API has already been paid for."""
    result = build("scan.pdf", make_invoice(), CONFIG)
    for field in sorted(KNOWN_COLUMNS):
        field_value(result, field)  # raises KeyError if unmapped


# --------------------------------------------------------------------------
# discovery
# --------------------------------------------------------------------------


def test_discover_walks_directories_and_skips_others(tmp_path):
    (tmp_path / "sub").mkdir()
    (tmp_path / "a.pdf").write_bytes(b"x")
    (tmp_path / "sub" / "b.JPG").write_bytes(b"x")
    (tmp_path / "notes.txt").write_text("ignore me")
    found = discover([tmp_path])
    assert [p.name for p in found] == ["a.pdf", "b.JPG"]


def test_named_unsupported_file_is_an_error(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("x")
    with pytest.raises(UnsupportedDocument, match="unsupported file type"):
        discover([path])


def test_discover_does_not_double_count(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"x")
    found = discover([tmp_path, tmp_path / "a.pdf"])
    assert len(found) == 1


# --------------------------------------------------------------------------
# export
# --------------------------------------------------------------------------


def build_results():
    clean = build("clean.pdf", make_invoice(), CONFIG)
    clean.gl_account = "6410"
    clean.property_code = "HTL-002"

    held = build("held.pdf", make_invoice(invoice_number=text(None)), CONFIG)
    return [clean, held]


def test_only_passing_invoices_reach_the_upload_sheet(tmp_path: Path):
    out = tmp_path / "out.xlsx"
    write(out, build_results(), CONFIG)

    book = load_workbook(out)
    main = book["Invoices"]
    headers = [c.value for c in main[1]]
    assert headers == list(DEFAULT_COLUMNS.values())
    assert main.max_row == 2  # header + the one clean invoice
    row = {h: main.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
    assert row["Vendor"] == "Sysco Corporation"
    assert row["Invoice Number"] == "INV-00123"
    assert row["GL Account"] == "6410"
    assert row["Amount"] == Decimal("1100.00")
    assert row["Property"] == "HTL-002"


def test_held_invoices_appear_with_their_reason(tmp_path: Path):
    out = tmp_path / "out.xlsx"
    write(out, build_results(), CONFIG)

    sheet = load_workbook(out)["Exceptions"]
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).value == "held.pdf"
    assert "missing-field" in sheet.cell(row=2, column=8).value


def test_include_review_puts_everything_on_the_main_sheet(tmp_path: Path):
    out = tmp_path / "out.xlsx"
    write(out, build_results(), CONFIG, include_review_in_main=True)
    assert load_workbook(out)["Invoices"].max_row == 3


def test_audit_sheet_records_the_source_text(tmp_path: Path):
    out = tmp_path / "out.xlsx"
    write(out, build_results(), CONFIG)
    sheet = load_workbook(out)["Audit"]
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    amounts = [r for r in rows if r[0] == "clean.pdf" and r[2] == "total amount"]
    assert amounts and amounts[0][4] == "1100.00"  # text on invoice
    assert amounts[0][5] == "high"
