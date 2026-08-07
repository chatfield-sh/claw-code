"""End-to-end wiring, with the API calls stubbed out.

These exercise everything except the model itself: discovery, both passes,
consensus, validation, GL coding, duplicate detection, the workbook, the
console report and the exit code.
"""

from __future__ import annotations

import textwrap

import pytest
from conftest import make_invoice, money, text
from openpyxl import load_workbook

from invoice_extract import cli
from invoice_extract.config import Config, GLConfig
from invoice_extract.extract import ExtractionError, Usage
from invoice_extract.ledger import Ledger
from invoice_extract.models import (
    ExtractedDocument,
    VerificationDocument,
    VerificationInvoice,
)
from invoice_extract.pipeline import run

GL_YAML = """
accounts:
  "6410": Food Cost
vendors:
  - name: "Sysco Corporation"
    account: "6410"
"""


def verification(**overrides) -> VerificationInvoice:
    fields = {
        "vendor_name": "Sysco Corporation",
        "invoice_number": "INV-00123",
        "invoice_date": "2026-07-15",
        "total_amount": "1100.00",
    }
    fields.update(overrides)
    return VerificationInvoice(**fields)  # type: ignore[arg-type]


class StubExtractor:
    """Stands in for the API. Keyed by file name so a batch can mix outcomes."""

    def __init__(self, reads, verifications=None, fail=()):
        self.reads = reads
        self.verifications = verifications or {}
        self.fail = set(fail)
        self.usage = Usage()
        self.seen: list[str] = []

    def read(self, path):
        self.seen.append(path.name)
        if path.name in self.fail:
            raise ExtractionError(f"{path.name}: stubbed failure")
        return ExtractedDocument(invoices=self.reads[path.name])

    def verify(self, path):
        return VerificationDocument(invoices=self.verifications.get(path.name, []))


@pytest.fixture
def gl(tmp_path):
    path = tmp_path / "gl.yaml"
    path.write_text(textwrap.dedent(GL_YAML))
    return GLConfig.load(path)


def scans(tmp_path, *names):
    paths = []
    for name in names:
        path = tmp_path / name
        path.write_bytes(b"%PDF-1.4 stub")
        paths.append(path)
    return paths


def go(paths, extractor, gl, config=None, ledger_path=None):
    config = config or Config()
    with Ledger(ledger_path) as ledger:
        return run(paths, config, gl, ledger, extractor)  # type: ignore[arg-type]


def test_clean_batch_is_ready_to_upload(tmp_path, gl):
    paths = scans(tmp_path, "a.pdf")
    extractor = StubExtractor(
        {"a.pdf": [make_invoice()]}, {"a.pdf": [verification()]}
    )
    report = go(paths, extractor, gl)

    assert len(report.auto) == 1
    assert not report.review
    assert not report.failures
    assert report.auto[0].gl_account == "6410"


def test_verification_disagreement_holds_the_invoice(tmp_path, gl):
    paths = scans(tmp_path, "a.pdf")
    extractor = StubExtractor(
        {"a.pdf": [make_invoice()]},
        {"a.pdf": [verification(total_amount="1700.00")]},
    )
    report = go(paths, extractor, gl)

    assert not report.auto
    assert len(report.review) == 1
    assert "verify-total" in report.review[0].review_reasons


def test_two_invoices_in_one_file_are_both_processed(tmp_path, gl):
    second = make_invoice(
        invoice_number=text("INV-00124"),
        subtotal=money("500.00"),
        tax=money("40.00"),
        freight=money("10.00"),
        total_amount=money("550.00"),
        line_items=[],
    )
    paths = scans(tmp_path, "batch.pdf")
    extractor = StubExtractor(
        {"batch.pdf": [make_invoice(), second]},
        {
            "batch.pdf": [
                verification(),
                verification(invoice_number="INV-00124", total_amount="550.00"),
            ]
        },
    )
    report = go(paths, extractor, gl)
    assert len(report.auto) == 2


def test_unreadable_file_is_reported_not_silently_dropped(tmp_path, gl):
    paths = scans(tmp_path, "good.pdf", "bad.pdf")
    extractor = StubExtractor(
        {"good.pdf": [make_invoice()]},
        {"good.pdf": [verification()]},
        fail={"bad.pdf"},
    )
    report = go(paths, extractor, gl)

    assert len(report.auto) == 1
    assert [f.path.name for f in report.failures] == ["bad.pdf"]
    assert report.files_read == 2


def test_file_with_no_invoice_is_reported(tmp_path, gl):
    paths = scans(tmp_path, "packing-slip.pdf")
    extractor = StubExtractor({"packing-slip.pdf": []})
    report = go(paths, extractor, gl)
    assert not report.results
    assert "no invoice" in report.failures[0].reason


def test_verification_returning_nothing_holds_the_invoice(tmp_path, gl):
    """A verification pass that yields no invoices must not read as agreement."""
    paths = scans(tmp_path, "a.pdf")
    extractor = StubExtractor({"a.pdf": [make_invoice()]}, {"a.pdf": []})
    report = go(paths, extractor, gl)
    assert "verify-count" in report.review[0].review_reasons


def test_no_verify_skips_the_second_pass(tmp_path, gl):
    paths = scans(tmp_path, "a.pdf")
    extractor = StubExtractor({"a.pdf": [make_invoice()]})
    report = go(paths, extractor, gl, config=Config(verify=False))
    assert len(report.auto) == 1
    assert report.auto[0].verification is None


def test_results_follow_input_order_regardless_of_concurrency(tmp_path, gl):
    names = [f"{n:02d}.pdf" for n in range(9)]
    paths = scans(tmp_path, *names)
    reads = {name: [make_invoice(invoice_number=text(f"INV-{name}"))] for name in names}
    checks = {
        name: [verification(invoice_number=f"INV-{name}")] for name in names
    }
    report = go(paths, StubExtractor(reads, checks), gl)
    assert [r.source_file for r in report.results] == names


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def test_cli_writes_the_workbook_and_signals_review(tmp_path, monkeypatch, capsys):
    gl_path = tmp_path / "gl.yaml"
    gl_path.write_text(textwrap.dedent(GL_YAML))
    scans(tmp_path, "clean.pdf", "held.pdf")

    extractor = StubExtractor(
        {
            "clean.pdf": [make_invoice()],
            "held.pdf": [make_invoice(invoice_number=text(None))],
        },
        {
            "clean.pdf": [verification()],
            "held.pdf": [verification(invoice_number=None)],
        },
    )
    monkeypatch.setattr(cli, "Extractor", lambda **kwargs: extractor)

    out = tmp_path / "upload.xlsx"
    code = cli.main(
        [
            str(tmp_path),
            "--gl",
            str(gl_path),
            "--out",
            str(out),
            "--ledger",
            str(tmp_path / "l.sqlite"),
        ]
    )

    assert code == cli.EXIT_REVIEW  # one invoice needs a human
    book = load_workbook(out)
    assert book["Invoices"].max_row == 2  # header + the clean one only
    assert book["Exceptions"].max_row == 2

    report = capsys.readouterr().err
    assert "1 invoice(s) ready to upload" in report
    assert "1 invoice(s) held for review" in report


def test_cli_returns_zero_when_everything_passes(tmp_path, monkeypatch):
    gl_path = tmp_path / "gl.yaml"
    gl_path.write_text(textwrap.dedent(GL_YAML))
    scans(tmp_path, "clean.pdf")
    extractor = StubExtractor(
        {"clean.pdf": [make_invoice()]}, {"clean.pdf": [verification()]}
    )
    monkeypatch.setattr(cli, "Extractor", lambda **kwargs: extractor)

    code = cli.main(
        [
            str(tmp_path),
            "--gl",
            str(gl_path),
            "--out",
            str(tmp_path / "o.xlsx"),
            "--ledger",
            str(tmp_path / "l.sqlite"),
        ]
    )
    assert code == cli.EXIT_OK


def test_cli_rerun_catches_the_resubmitted_batch(tmp_path, monkeypatch):
    gl_path = tmp_path / "gl.yaml"
    gl_path.write_text(textwrap.dedent(GL_YAML))
    scans(tmp_path, "clean.pdf")
    ledger = tmp_path / "l.sqlite"

    def invoke():
        extractor = StubExtractor(
            {"clean.pdf": [make_invoice()]}, {"clean.pdf": [verification()]}
        )
        monkeypatch.setattr(cli, "Extractor", lambda **kwargs: extractor)
        return cli.main(
            [
                str(tmp_path / "clean.pdf"),
                "--gl",
                str(gl_path),
                "--out",
                str(tmp_path / "o.xlsx"),
                "--ledger",
                str(ledger),
            ]
        )

    assert invoke() == cli.EXIT_OK
    assert invoke() == cli.EXIT_REVIEW  # same invoice, second time


def test_cli_list_stops_before_calling_the_api(tmp_path, capsys):
    scans(tmp_path, "a.pdf", "b.pdf")
    assert cli.main([str(tmp_path), "--list"]) == cli.EXIT_OK
    assert "2 file(s)" in capsys.readouterr().out


def test_cli_rejects_a_bad_gl_file(tmp_path, capsys):
    gl_path = tmp_path / "gl.yaml"
    gl_path.write_text('accounts:\n  "1": x\nvendors:\n  - name: A\n    account: "2"\n')
    scans(tmp_path, "a.pdf")
    assert cli.main([str(tmp_path), "--gl", str(gl_path)]) == cli.EXIT_ERROR
    assert "not in `accounts`" in capsys.readouterr().err
